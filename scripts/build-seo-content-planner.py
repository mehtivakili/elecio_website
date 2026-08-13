from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUT = Path("deliverables/Elecio_SEO_Content_Planner_FA.xlsx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "0A1224"
BLUE = "1473E6"
CYAN = "25C2FF"
PALE_BLUE = "EAF4FF"
PALE_CYAN = "E9FBFF"
WHITE = "FFFFFF"
INK = "172033"
MUTED = "64748B"
GRID = "D7E0EC"
GREEN = "1F9D68"
PALE_GREEN = "E8F7F0"
AMBER = "D88700"
PALE_AMBER = "FFF5D9"
RED = "C43D4D"
PALE_RED = "FDECEF"
PURPLE = "7357C8"

thin = Side(style="thin", color=GRID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def fa_font(size=10, bold=False, color=INK):
    return Font(name="Vazirmatn", size=size, bold=bold, color=color)


def setup_sheet(ws, freeze="A2", zoom=90):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    ws.sheet_view.zoomScale = zoom
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25


def title_band(ws, title, subtitle, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    ws.cell(1, 1, title).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).font = fa_font(18, True, WHITE)
    ws.cell(1, 1).alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 38
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    ws.cell(2, 1, subtitle).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(2, 1).font = fa_font(10, False, MUTED)
    ws.cell(2, 1).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 32


def header_row(ws, row, headers):
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = fa_font(10, True, WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 34


def style_data(ws, min_row, max_row, min_col, max_col, height=42):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.font = fa_font()
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = border
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFD")
        ws.row_dimensions[row[0].row].height = height


def add_table(ws, ref, name):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(tab)


def add_list_validation(ws, cell_range, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "یکی از گزینه‌های فهرست را انتخاب کنید."
    dv.errorTitle = "مقدار نامعتبر"
    dv.prompt = "از فهرست انتخاب کنید."
    dv.promptTitle = "انتخاب مقدار"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def safe_slug(text):
    replacements = {
        " ": "-", "‌": "-", "؟": "", "،": "", "/": "-", "(": "", ")": "",
        "آ": "a", "ا": "a", "ب": "b", "پ": "p", "ت": "t", "ث": "s", "ج": "j",
        "چ": "ch", "ح": "h", "خ": "kh", "د": "d", "ذ": "z", "ر": "r", "ز": "z",
        "ژ": "zh", "س": "s", "ش": "sh", "ص": "s", "ض": "z", "ط": "t", "ظ": "z",
        "ع": "a", "غ": "gh", "ف": "f", "ق": "gh", "ک": "k", "گ": "g", "ل": "l",
        "م": "m", "ن": "n", "و": "v", "ه": "h", "ی": "y", "ئ": "y", "ء": "",
    }
    result = "".join(replacements.get(ch, ch.lower()) for ch in text)
    while "--" in result:
        result = result.replace("--", "-")
    return result.strip("-")[:72]


def standard_h2(keyword, mode):
    templates = {
        "راهنمای جامع": [
            f"{keyword} چیست و چه مسئله‌ای را حل می‌کند؟",
            "اجزای اصلی و پیش‌نیازهای فنی",
            "مراحل راه‌اندازی و استفاده درست",
            "اشتباهات رایج و روش پیشگیری",
            "چک‌لیست انتخاب یا اجرای مطمئن",
            "سؤالات متداول",
        ],
        "آموزش فنی": [
            f"پیش‌نیازهای {keyword}",
            "تنظیمات لازم در کنترلر و سمت شبکه",
            "مراحل اجرا به‌صورت گام‌به‌گام",
            "خطاهای متداول و روش عیب‌یابی",
            "نکات امنیتی و نگهداری",
            "سؤالات متداول",
        ],
        "راهنمای خرید": [
            f"قبل از انتخاب {keyword} چه چیزهایی را مشخص کنیم؟",
            "معیارهای فنی مؤثر در انتخاب",
            "سازگاری با کنترلرها و زیرساخت کارگاه",
            "هزینه کل مالکیت و بازگشت سرمایه",
            "چک‌لیست دریافت مشاوره و استعلام",
            "سؤالات متداول",
        ],
        "مقایسه": [
            f"تعریف گزینه‌های مورد مقایسه در {keyword}",
            "مقایسه سرعت، امنیت و قابلیت اطمینان",
            "مقایسه هزینه اجرا و نگهداری",
            "هر روش برای چه نوع کارگاهی مناسب است؟",
            "جدول تصمیم‌گیری نهایی",
            "سؤالات متداول",
        ],
        "عیب‌یابی": [
            f"نشانه‌های اصلی مشکل {keyword}",
            "بررسی کابل، پورت، شبکه و تنظیمات",
            "خطاهای رایج کنترلر و نرم‌افزار",
            "مراحل عیب‌یابی از ساده به تخصصی",
            "چه زمانی به پشتیبانی فنی نیاز است؟",
            "سؤالات متداول",
        ],
        "کاربرد صنعتی": [
            f"چرا {keyword} برای کارگاه CNC مهم است؟",
            "سناریوی واقعی قبل و بعد از اجرا",
            "اثر بر زمان، خطا و دسترس‌پذیری تولید",
            "الزامات اجرا در مقیاس کوچک و بزرگ",
            "شاخص‌های سنجش نتیجه",
            "سؤالات متداول",
        ],
    }
    return "\n".join(templates.get(mode, templates["راهنمای جامع"]))


def topic(cluster, keyword, title, secondary, intent="اطلاعاتی", funnel="آگاهی", mode="راهنمای جامع",
          target="/fa/product", cta="دریافت مشاوره فنی", status="ایده", priority="متوسط", source="کاتالوگ + تجربه فنی ElecIO"):
    seo_title = title.replace("؟", "").replace("؛", " |")
    if len(seo_title) > 58:
        seo_title = f"{keyword} | راهنمای تخصصی ElecIO"
    meta = f"در این راهنمای تخصصی، {keyword} را با تمرکز بر کاربرد واقعی در کارگاه CNC، نکات فنی، خطاهای رایج و روش انتخاب یا اجرای مطمئن بررسی می‌کنیم."
    return {
        "cluster": cluster, "keyword": keyword, "title": title, "secondary": secondary,
        "intent": intent, "funnel": funnel, "mode": mode, "target": target, "cta": cta,
        "status": status, "priority": priority, "source": source, "seo": seo_title,
        "meta": meta, "slug": safe_slug(keyword), "h2": standard_h2(keyword, mode),
    }


topics = [
    topic("انتقال فایل و DNC", "انتقال فایل CNC", "راهنمای کامل انتقال فایل CNC بدون فلش و رفت‌وآمد", "انتقال برنامه CNC، ارسال فایل NC، شبکه CNC", mode="آموزش فنی", status="منتشرشده", priority="بالا"),
    topic("کنترلر FANUC", "تنظیم RS232 فانوک", "چک‌لیست تنظیم RS232 فانوک برای انتقال مطمئن برنامه", "پارامتر RS232 فانوک، baud rate فانوک، کابل سریال CNC", mode="آموزش فنی", status="منتشرشده", priority="بالا"),
    topic("کنترلر Siemens", "انتقال برنامه زیمنس", "انتقال برنامه زیمنس به CNC؛ مسیر شبکه و نکات سازگاری", "Siemens CNC network، انتقال فایل Sinumerik، DNC زیمنس", mode="آموزش فنی", status="منتشرشده", priority="بالا"),
    topic("امنیت و پشتیبان‌گیری", "بکاپ برنامه CNC", "بکاپ برنامه CNC؛ اصول امنیت، نسخه‌بندی و بازیابی", "پشتیبان گیری CNC، آرشیو G-code، امنیت فایل CNC", mode="راهنمای جامع", status="منتشرشده", priority="بالا"),
    topic("گزارش و بهره‌وری", "گزارش توقف دستگاه CNC", "گزارش توقف دستگاه CNC؛ تبدیل داده خام به تصمیم تولید", "دلایل توقف CNC، گزارش تولید، زمان کارکرد دستگاه", mode="کاربرد صنعتی", status="منتشرشده", priority="بالا", target="/fa/software"),
    topic("انتقال فایل و DNC", "DNC چیست", "DNC چیست و چرا کارگاه‌های CNC به آن نیاز دارند؟", "سیستم DNC، نرم افزار DNC، شبکه دستگاه CNC", priority="بالا"),
    topic("انتقال فایل و DNC", "دستگاه DNC", "دستگاه DNC چگونه انتقال برنامه CNC را ساده می‌کند؟", "سخت افزار DNC، انتقال فایل بی سیم CNC، DNC صنعتی", intent="تجاری", funnel="بررسی", mode="راهنمای خرید", priority="بالا"),
    topic("انتقال فایل و DNC", "نرم افزار DNC", "نرم افزار DNC چه امکاناتی برای مدیریت برنامه‌های CNC لازم دارد؟", "مدیریت فایل CNC، نرم افزار انتقال برنامه، DNC server", intent="تجاری", funnel="بررسی", mode="راهنمای خرید", target="/fa/software", priority="بالا"),
    topic("انتقال فایل و DNC", "قیمت دستگاه DNC", "قیمت دستگاه DNC به چه عواملی بستگی دارد؟", "هزینه DNC، خرید DNC، قیمت انتقال فایل CNC", intent="خرید", funnel="تصمیم", mode="راهنمای خرید", cta="استعلام قیمت Smart DNC", priority="بالا"),
    topic("انتقال فایل و DNC", "خرید DNC", "راهنمای خرید DNC برای کارگاه و کارخانه CNC", "بهترین DNC، انتخاب DNC، DNC برای چند دستگاه", intent="خرید", funnel="تصمیم", mode="راهنمای خرید", cta="مشاوره انتخاب Smart DNC", priority="بالا"),
    topic("انتقال فایل و DNC", "انتقال فایل CNC بدون فلش", "چرا انتقال فایل CNC بدون فلش امن‌تر و سریع‌تر است؟", "حذف فلش از CNC، انتقال وایرلس CNC، ویروس CNC", intent="اطلاعاتی/تجاری", funnel="بررسی", mode="مقایسه", priority="بالا"),
    topic("شبکه صنعتی", "انتقال وایرلس CNC", "انتقال وایرلس CNC؛ مزایا، محدودیت‌ها و روش اجرا", "Wi-Fi CNC، ارسال بی سیم فایل CNC، شبکه کارگاه", mode="راهنمای جامع", priority="بالا"),
    topic("شبکه صنعتی", "شبکه CNC", "شبکه CNC چیست و چگونه چند دستگاه را متصل می‌کند؟", "اتصال CNC به شبکه، شبکه صنعتی، DNC network", mode="راهنمای جامع", priority="بالا"),
    topic("شبکه صنعتی", "اتصال CNC به شبکه", "اتصال CNC به شبکه؛ راهنمای طراحی از کنترلر تا سرور", "LAN CNC، پورت شبکه CNC، انتقال برنامه تحت شبکه", mode="آموزش فنی", priority="بالا"),
    topic("شبکه صنعتی", "امنیت شبکه CNC", "امنیت شبکه CNC؛ کنترل دسترسی بدون توقف تولید", "امنیت صنعتی، جداسازی شبکه CNC، دسترسی فایل NC", mode="راهنمای جامع", priority="بالا", target="/fa/software"),
    topic("امنیت و پشتیبان‌گیری", "ویروس فلش CNC", "ویروس فلش CNC چه ریسکی برای برنامه و تولید ایجاد می‌کند؟", "بدافزار CNC، حذف USB از کارگاه، امنیت فایل G-code", mode="کاربرد صنعتی", priority="بالا"),
    topic("ارتباط سریال", "پورت RS232 CNC", "پورت RS232 در CNC چیست و چگونه کار می‌کند؟", "پورت سریال CNC، ارتباط سریال، DB25 CNC", mode="راهنمای جامع"),
    topic("ارتباط سریال", "کابل RS232 CNC", "راهنمای انتخاب و سیم‌بندی کابل RS232 برای CNC", "پین اوت RS232، کابل DB9 DB25، شیلد کابل سریال", mode="آموزش فنی", priority="بالا"),
    topic("ارتباط سریال", "خطای RS232 CNC", "خطای RS232 در CNC؛ چک‌لیست عیب‌یابی ارتباط سریال", "خطای انتقال فایل CNC، parity error، baud rate", intent="پشتیبانی", funnel="پس از خرید", mode="عیب‌یابی", target="/fa/contact", cta="درخواست بررسی فنی", priority="بالا"),
    topic("کنترلر FANUC", "baud rate فانوک", "تنظیم Baud Rate فانوک برای انتقال بدون خطای برنامه", "پارامتر I/O فانوک، parity فانوک، RS232 Fanuc", mode="آموزش فنی"),
    topic("کنترلر FANUC", "پارامترهای انتقال فایل فانوک", "پارامترهای انتقال فایل فانوک؛ چک‌لیست قبل از اتصال", "I/O channel Fanuc، تنظیم سریال فانوک، PUNCH READ", mode="آموزش فنی", priority="بالا"),
    topic("کنترلر FANUC", "انتقال برنامه فانوک", "انتقال برنامه فانوک به کامپیوتر و برعکس", "ارسال برنامه Fanuc، دریافت فایل CNC، DNC Fanuc", mode="آموزش فنی", priority="بالا"),
    topic("کنترلر Heidenhain", "انتقال برنامه هایدن هاین", "انتقال برنامه Heidenhain؛ از TNC تا شبکه کارگاه", "TNC server، انتقال فایل Heidenhain، LSV2", mode="آموزش فنی", priority="بالا"),
    topic("کنترلر Heidenhain", "اتصال Heidenhain به شبکه", "اتصال کنترلر Heidenhain به شبکه؛ نکات سازگاری و امنیت", "شبکه TNC، LSV2، TNCremo", mode="آموزش فنی"),
    topic("سازگاری کنترلر", "انتقال برنامه Mitsubishi CNC", "انتقال برنامه Mitsubishi CNC؛ روش‌ها و خطاهای رایج", "DNC Mitsubishi، RS232 میتسوبیشی، ارسال G-code", mode="آموزش فنی"),
    topic("سازگاری کنترلر", "انتقال برنامه Haas", "انتقال برنامه Haas؛ انتخاب روش سریال یا شبکه", "DNC Haas، RS232 Haas، انتقال فایل HAAS", mode="مقایسه"),
    topic("سازگاری کنترلر", "انتقال برنامه Charmilles", "انتقال برنامه Charmilles؛ نکات ارتباط با دستگاه EDM", "DNC Charmilles، ارسال فایل EDM، ارتباط سریال", mode="آموزش فنی"),
    topic("امنیت و پشتیبان‌گیری", "پشتیبان گیری CNC", "پشتیبان‌گیری CNC؛ چه فایل‌هایی را و هر چند وقت ذخیره کنیم؟", "بکاپ پارامتر CNC، بکاپ برنامه، بازیابی CNC", mode="راهنمای جامع", priority="بالا"),
    topic("مدیریت برنامه", "مدیریت فایل CNC", "مدیریت فایل CNC؛ نام‌گذاری، جست‌وجو و جلوگیری از نسخه اشتباه", "آرشیو فایل NC، مدیریت G-code، کنترل نسخه CNC", mode="راهنمای جامع", target="/fa/software", priority="بالا"),
    topic("مدیریت برنامه", "نسخه بندی برنامه CNC", "نسخه‌بندی برنامه CNC؛ پایان فایل‌های تکراری و نامشخص", "revision برنامه CNC، تاریخچه فایل NC، کنترل تغییرات", mode="کاربرد صنعتی", target="/fa/software"),
    topic("نرم‌افزار", "نرم افزار مدیریت دستگاه CNC", "نرم افزار مدیریت دستگاه CNC باید چه امکاناتی داشته باشد؟", "مدیریت تولید CNC، مانیتورینگ CNC، نرم افزار DNC", intent="تجاری", funnel="بررسی", mode="راهنمای خرید", target="/fa/software", priority="بالا"),
    topic("نرم‌افزار", "مانیتورینگ CNC", "مانیتورینگ CNC چیست و چه داده‌ای برای مدیر تولید ارزش دارد؟", "پایش دستگاه CNC، وضعیت ماشین، داشبورد تولید", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("گزارش و بهره‌وری", "گزارش تولید CNC", "گزارش تولید CNC؛ شاخص‌هایی که باید هر روز ببینید", "گزارش کارکرد دستگاه، توقف تولید، داشبورد CNC", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("گزارش و بهره‌وری", "OEE دستگاه CNC", "OEE دستگاه CNC چیست و چگونه درست محاسبه می‌شود؟", "اثربخشی کلی تجهیزات، بهره وری CNC، availability CNC", mode="راهنمای جامع", target="/fa/software"),
    topic("گزارش و بهره‌وری", "کاهش توقف دستگاه CNC", "کاهش توقف دستگاه CNC با داده‌های واقعی تولید", "کاهش downtime، علت توقف CNC، افزایش بهره وری", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("گزارش و بهره‌وری", "علت توقف دستگاه CNC", "علت توقف دستگاه CNC؛ دسته‌بندی و تحلیل ریشه‌ای", "downtime CNC، خرابی CNC، توقف برنامه‌ریزی نشده", mode="عیب‌یابی", target="/fa/software"),
    topic("گزارش و بهره‌وری", "زمان کارکرد دستگاه CNC", "زمان کارکرد دستگاه CNC را چگونه دقیق ثبت کنیم؟", "runtime CNC، ساعت کار ماشین، گزارش اپراتور", mode="کاربرد صنعتی", target="/fa/software"),
    topic("تحول دیجیتال", "هوشمندسازی کارگاه CNC", "هوشمندسازی کارگاه CNC را از کجا شروع کنیم؟", "دیجیتال سازی CNC، صنعت 4، اتصال ماشین آلات", mode="کاربرد صنعتی", priority="بالا"),
    topic("تحول دیجیتال", "صنعت 4 در CNC", "صنعت 4 در CNC؛ از اتصال دستگاه تا تصمیم داده‌محور", "Industry 4.0 CNC، کارخانه هوشمند، داده تولید", mode="کاربرد صنعتی"),
    topic("تحول دیجیتال", "اینترنت اشیا صنعتی CNC", "اینترنت اشیا صنعتی در CNC؛ کاربرد واقعی بدون پیچیدگی اضافی", "IIoT CNC، اتصال ماشین قدیمی، داده صنعتی", mode="راهنمای جامع"),
    topic("مقایسه و انتخاب", "مقایسه DNC و فلش", "مقایسه DNC و فلش USB برای انتقال برنامه CNC", "انتقال فایل CNC، امنیت USB، سرعت انتقال برنامه", intent="تجاری", funnel="بررسی", mode="مقایسه", priority="بالا"),
    topic("مقایسه و انتخاب", "DNC و Drip Feed", "تفاوت DNC و Drip Feed چیست و کدام را نیاز دارید؟", "drip feed CNC، اجرای برنامه بزرگ، حافظه CNC", mode="مقایسه"),
    topic("انتقال فایل و DNC", "Drip Feed CNC", "Drip Feed در CNC چیست؟ راهنمای اجرای برنامه‌های بزرگ", "DNC mode، حافظه کنترلر، اجرای آنلاین G-code", mode="راهنمای جامع"),
    topic("آموزش پایه", "فایل NC چیست", "فایل NC چیست و چه تفاوتی با G-code دارد؟", "پسوند فایل CNC، برنامه NC، ساختار G-code", mode="راهنمای جامع", target="/fa/blog"),
    topic("آموزش پایه", "انتقال G-code به CNC", "انتقال G-code به CNC؛ روش‌ها، کنترل صحت و اجرای امن", "ارسال G-code، فایل NC، انتقال برنامه", mode="آموزش فنی", priority="بالا"),
    topic("عیب‌یابی", "خطای ارسال فایل به CNC", "خطای ارسال فایل به CNC؛ از کابل تا پارامتر کنترلر", "خطای DNC، خطای سریال CNC، فایل ناقص", intent="پشتیبانی", funnel="پس از خرید", mode="عیب‌یابی", target="/fa/contact", cta="ارسال مشخصات خطا", priority="بالا"),
    topic("نصب و راه‌اندازی", "نصب DNC", "چک‌لیست نصب DNC کنار دستگاه CNC", "راه اندازی DNC، کابل DNC، تنظیم شبکه CNC", mode="آموزش فنی", priority="بالا"),
    topic("امنیت و پشتیبان‌گیری", "امنیت G-code", "امنیت G-code؛ جلوگیری از اجرای نسخه اشتباه یا دستکاری‌شده", "کنترل دسترسی CNC، امنیت برنامه NC، تایید فایل", mode="راهنمای جامع", target="/fa/software"),
    topic("مدیریت برنامه", "مدیریت چند دستگاه CNC", "مدیریت چند دستگاه CNC از یک نقطه؛ معماری و چالش‌ها", "DNC مرکزی، شبکه چند CNC، مدیریت فایل کارخانه", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("شبکه صنعتی", "اتصال CNC قدیمی به شبکه", "اتصال CNC قدیمی به شبکه بدون تعویض کنترلر", "retrofit CNC، RS232 to network، هوشمندسازی ماشین قدیمی", intent="تجاری", funnel="بررسی", mode="راهنمای خرید", priority="بالا"),
    topic("انتقال فایل و DNC", "DNC برای کارگاه کوچک", "آیا DNC برای کارگاه کوچک CNC به‌صرفه است؟", "DNC تک دستگاه، هزینه انتقال فایل، کارگاه CNC", intent="تجاری", funnel="بررسی", mode="کاربرد صنعتی", priority="بالا"),
    topic("انتقال فایل و DNC", "DNC برای کارخانه", "طراحی DNC برای کارخانه؛ مقیاس، دسترسی و پایداری", "DNC مرکزی، شبکه کارخانه CNC، مدیریت کاربران", intent="تجاری", funnel="بررسی", mode="راهنمای خرید"),
    topic("اقتصاد تولید", "هزینه توقف دستگاه CNC", "هزینه توقف دستگاه CNC را چگونه محاسبه کنیم؟", "هزینه downtime، زیان توقف تولید، بهره وری ماشین", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("اقتصاد تولید", "بازگشت سرمایه DNC", "بازگشت سرمایه DNC؛ محاسبه ROI برای کارگاه CNC", "ROI هوشمندسازی، هزینه DNC، کاهش توقف", intent="تجاری", funnel="تصمیم", mode="کاربرد صنعتی", cta="محاسبه سناریوی کارگاه شما", priority="بالا"),
    topic("مدیریت برنامه", "ردیابی تغییرات برنامه CNC", "ردیابی تغییرات برنامه CNC؛ چه کسی چه چیزی را عوض کرد؟", "audit trail CNC، تاریخچه برنامه، کنترل نسخه", mode="کاربرد صنعتی", target="/fa/software"),
    topic("امنیت و پشتیبان‌گیری", "کنترل دسترسی فایل CNC", "کنترل دسترسی فایل CNC؛ نقش‌ها، مجوزها و ثبت رویداد", "سطح دسترسی CNC، امنیت فایل NC، مدیریت کاربران", mode="راهنمای جامع", target="/fa/software"),
    topic("نصب و راه‌اندازی", "راه اندازی Smart DNC", "راه‌اندازی Smart DNC؛ از اتصال دستگاه تا اولین انتقال فایل", "نصب ElecIO، تنظیم Smart DNC، اتصال CNC", intent="پشتیبانی", funnel="پس از خرید", mode="آموزش فنی", target="/fa/product", cta="دریافت راهنمای راه‌اندازی", priority="بالا"),
    topic("نرم‌افزار", "گزارش گیری دستگاه CNC", "گزارش‌گیری دستگاه CNC با ElecIO؛ چه چیزی قابل مشاهده است؟", "گزارش ElecIO، گزارش زمان کار، وضعیت دستگاه", intent="تجاری", funnel="بررسی", mode="کاربرد صنعتی", target="/fa/software", priority="بالا"),
    topic("مقایسه و انتخاب", "بهترین روش انتقال فایل CNC", "بهترین روش انتقال فایل CNC برای کنترلرهای قدیمی و جدید", "USB CNC، RS232 CNC، Ethernet CNC، DNC", intent="تجاری", funnel="بررسی", mode="مقایسه", priority="بالا"),
    topic("اعتماد و برند", "ElecIO Smart DNC", "ElecIO Smart DNC چیست و برای چه کارگاهی مناسب است؟", "الکسیو DNC، دستگاه انتقال فایل CNC، Smart DNC", intent="برند/خرید", funnel="تصمیم", mode="راهنمای خرید", cta="سفارش Smart DNC", priority="بالا", source="کاتالوگ رسمی ElecIO + صفحه محصول"),
]


existing_urls = {
    1: "/fa/blog/cnc-file-transfer-guide",
    2: "/fa/blog/fanuc-rs232-checklist",
    3: "/fa/blog/siemens-program-transfer-workflow",
    4: "/fa/blog/cnc-backup-security-basics",
    5: "/fa/blog/cnc-downtime-reporting-guide",
}

wb = Workbook()
wb.remove(wb.active)

# Hidden lists for validation
lists = wb.create_sheet("_lists")
validation_lists = {
    "A": ["ایده", "در انتظار بریف", "در حال تحقیق", "در حال نگارش", "در حال ویرایش", "آماده انتشار", "زمان‌بندی‌شده", "منتشرشده", "نیازمند به‌روزرسانی"],
    "B": ["خیلی بالا", "بالا", "متوسط", "پایین"],
    "C": ["اطلاعاتی", "تجاری", "خرید", "پشتیبانی", "برند/خرید", "اطلاعاتی/تجاری"],
    "D": ["آگاهی", "بررسی", "تصمیم", "پس از خرید"],
    "E": ["بله", "خیر", "نامشخص"],
    "F": ["تأیید", "نیازمند اصلاح", "بررسی نشده", "نامرتبط"],
    "G": ["راهنمای جامع", "آموزش فنی", "راهنمای خرید", "مقایسه", "عیب‌یابی", "کاربرد صنعتی"],
}
for col, values in validation_lists.items():
    for i, value in enumerate(values, 1):
        lists[f"{col}{i}"] = value
lists.sheet_state = "hidden"

# Dashboard
ws = wb.create_sheet("داشبورد")
setup_sheet(ws, freeze="A4", zoom=95)
title_band(ws, "داشبورد برنامه محتوای سئو ElecIO", "نمای کلی برنامه، پیشرفت انتشار و تمرکز خوشه‌های محتوایی — اعداد رتبه و حجم جست‌وجو باید از ابزار واقعی وارد شوند.", 10)
for col, width in {"A": 3, "B": 22, "C": 18, "D": 4, "E": 22, "F": 18, "G": 4, "H": 22, "I": 18, "J": 3}.items():
    ws.column_dimensions[col].width = width
cards = [
    ("B4", "C4", "کل موضوعات", "=COUNTA('تقویم محتوا'!A4:A200)"),
    ("E4", "F4", "منتشرشده", '=COUNTIF(\'تقویم محتوا\'!Z4:Z200,"منتشرشده")'),
    ("H4", "I4", "آماده/زمان‌بندی", '=COUNTIF(\'تقویم محتوا\'!Z4:Z200,"آماده انتشار")+COUNTIF(\'تقویم محتوا\'!Z4:Z200,"زمان‌بندی‌شده")'),
    ("B7", "C7", "اولویت بالا", '=COUNTIF(\'تقویم محتوا\'!AA4:AA200,"بالا")+COUNTIF(\'تقویم محتوا\'!AA4:AA200,"خیلی بالا")'),
    ("E7", "F7", "میانگین پیشرفت", '=IFERROR(AVERAGE(\'تقویم محتوا\'!AB4:AB200),0)'),
    ("H7", "I7", "کیوردهای ثبت‌شده", "=COUNTA('نقشه کیورد'!B4:B250)"),
]
for label_cell, value_cell, label, formula in cards:
    ws[label_cell] = label
    ws[label_cell].fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws[label_cell].font = fa_font(11, True, MUTED)
    ws[label_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws[label_cell].border = border
    ws[value_cell] = formula
    ws[value_cell].fill = PatternFill("solid", fgColor=WHITE)
    ws[value_cell].font = fa_font(18, True, BLUE)
    ws[value_cell].alignment = Alignment(horizontal="center", vertical="center")
    ws[value_cell].border = border
    ws.row_dimensions[ws[label_cell].row].height = 48
ws["F7"].number_format = "0%"

ws["B11"] = "وضعیت"
ws["C11"] = "تعداد"
for c in (ws["B11"], ws["C11"]):
    c.fill = PatternFill("solid", fgColor=BLUE); c.font = fa_font(10, True, WHITE); c.border = border
statuses = validation_lists["A"]
for idx, status in enumerate(statuses, 12):
    ws.cell(idx, 2, status)
    ws.cell(idx, 3, f'=COUNTIF(\'تقویم محتوا\'!$Z$4:$Z$200,B{idx})')
    for c in ws[idx][1:3]: c.border = border; c.font = fa_font(); c.fill = PatternFill("solid", fgColor="F8FAFD")
chart = DoughnutChart()
chart.title = "توزیع وضعیت محتوا"
chart.add_data(Reference(ws, min_col=3, min_row=11, max_row=11 + len(statuses)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=2, min_row=12, max_row=11 + len(statuses)))
chart.height = 8
chart.width = 13
chart.legend.position = "r"
ws.add_chart(chart, "E11")

clusters = sorted({t["cluster"] for t in topics})
ws["B23"] = "خوشه"
ws["C23"] = "تعداد"
for c in (ws["B23"], ws["C23"]):
    c.fill = PatternFill("solid", fgColor=BLUE); c.font = fa_font(10, True, WHITE); c.border = border
for idx, cluster in enumerate(clusters, 24):
    ws.cell(idx, 2, cluster)
    ws.cell(idx, 3, f'=COUNTIF(\'تقویم محتوا\'!$D$4:$D$200,B{idx})')
    for c in ws[idx][1:3]: c.border = border; c.font = fa_font(); c.fill = PatternFill("solid", fgColor="F8FAFD")
bar = BarChart()
bar.type = "bar"
bar.style = 10
bar.title = "پوشش خوشه‌های موضوعی"
bar.add_data(Reference(ws, min_col=3, min_row=23, max_row=23 + len(clusters)), titles_from_data=True)
bar.set_categories(Reference(ws, min_col=2, min_row=24, max_row=23 + len(clusters)))
bar.height = 11
bar.width = 17
bar.legend = None
ws.add_chart(bar, "E23")

# Keyword map
ws = wb.create_sheet("نقشه کیورد")
setup_sheet(ws, freeze="A4", zoom=80)
headers = ["شناسه", "کلمه کلیدی اصلی", "خوشه موضوعی", "کلمات مرتبط", "نیت جست‌وجو", "مرحله قیف", "نوع محتوای پیشنهادی", "صفحه هدف/پیلار", "اولویت", "منبع داده", "حجم جست‌وجوی ماهانه", "سختی کیورد", "رتبه فعلی", "رتبه هدف", "کلیک GSC", "ایمپرشن GSC", "CTR", "ریسک هم‌نوع‌خواری", "شناسه محتوا", "آخرین بررسی", "یادداشت"]
title_band(ws, "نقشه کلمات کلیدی ElecIO", "حجم، سختی، رتبه و داده‌های Search Console خالی‌اند تا فقط از منابع واقعی پر شوند؛ ستون‌های استراتژیک از موضوعات سایت آماده شده‌اند.", len(headers))
header_row(ws, 3, headers)
for i, t in enumerate(topics, 1):
    row = 3 + i
    values = [f"KW-{i:03d}", t["keyword"], t["cluster"], t["secondary"], t["intent"], t["funnel"], t["mode"], t["target"], t["priority"], "نیاز به تحقیق در GSC / Keyword Planner", None, None, None, 10, None, None, None, "نامشخص", f"EC-{i:03d}", None, ""]
    for col, value in enumerate(values, 1): ws.cell(row, col, value)
style_data(ws, 4, 3 + len(topics), 1, len(headers), 48)
add_table(ws, f"A3:U{3 + len(topics)}", "KeywordMap")
for col, width in enumerate([13, 28, 24, 48, 18, 14, 22, 28, 12, 32, 16, 14, 13, 13, 12, 14, 12, 18, 14, 16, 35], 1): ws.column_dimensions[chr(64+col)].width = width
add_list_validation(ws, f"E4:E{3+len(topics)}", "'_lists'!$C$1:$C$6")
add_list_validation(ws, f"F4:F{3+len(topics)}", "'_lists'!$D$1:$D$4")
add_list_validation(ws, f"G4:G{3+len(topics)}", "'_lists'!$G$1:$G$6")
add_list_validation(ws, f"I4:I{3+len(topics)}", "'_lists'!$B$1:$B$4")
add_list_validation(ws, f"R4:R{3+len(topics)}", "'_lists'!$E$1:$E$3")
ws.conditional_formatting.add(f"I4:I{3+len(topics)}", FormulaRule(formula=['I4="بالا"'], fill=PatternFill("solid", fgColor=PALE_AMBER)))
ws.conditional_formatting.add(f"M4:M{3+len(topics)}", CellIsRule(operator="greaterThan", formula=[20], fill=PatternFill("solid", fgColor=PALE_RED)))

# Calendar
ws = wb.create_sheet("تقویم محتوا")
setup_sheet(ws, freeze="A4", zoom=70)
headers = ["شناسه", "هفته", "روز انتشار", "خوشه", "عنوان مقاله", "کیورد اصلی", "کیوردهای فرعی", "نیت", "مرحله قیف", "فرمت", "SEO Title", "طول Title", "Meta Description", "طول Meta", "Slug پیشنهادی", "H1", "H2های پیشنهادی", "لینک‌های داخلی", "منبع ضروری", "CTA", "هدف محتوا", "مسئول", "تاریخ بریف", "موعد انتشار", "لینک انتشار", "وضعیت", "اولویت", "درصد پیشرفت", "آخرین به‌روزرسانی", "یادداشت ویراستار"]
title_band(ws, "تقویم اجرایی محتوای سئو", "۶۰ موضوع اولویت‌بندی‌شده برای ۱۲ هفته؛ تاریخ‌ها و مسئول‌ها قابل ویرایش‌اند. پنج مقاله موجود در ابتدای برنامه ثبت شده‌اند.", len(headers))
header_row(ws, 3, headers)
days = ["شنبه", "یکشنبه", "دوشنبه", "سه‌شنبه", "چهارشنبه"]
start = date.today() + timedelta(days=(5 - date.today().weekday()) % 7)
for i, t in enumerate(topics, 1):
    row = 3 + i
    week = ((i - 1) // 5) + 1
    day = days[(i - 1) % 5]
    publish_date = start + timedelta(weeks=week - 1, days=(i - 1) % 5)
    internal = f"{t['target']} | صفحه محصول: /fa/product | تماس: /fa/contact"
    published_url = existing_urls.get(i, "")
    progress = 1 if t["status"] == "منتشرشده" else 0
    values = [f"EC-{i:03d}", f"هفته {week}", day, t["cluster"], t["title"], t["keyword"], t["secondary"], t["intent"], t["funnel"], t["mode"], t["seo"], f"=LEN(K{row})", t["meta"], f"=LEN(M{row})", t["slug"], t["title"], t["h2"], internal, t["source"], t["cta"], f"پوشش نیت «{t['intent']}» و هدایت کاربر به {t['target']}", "", "", publish_date, published_url, t["status"], t["priority"], progress, date.today() if t["status"] == "منتشرشده" else "", ""]
    for col, value in enumerate(values, 1): ws.cell(row, col, value)
    ws.cell(row, 24).number_format = "yyyy/mm/dd"
    ws.cell(row, 28).number_format = "0%"
    if published_url:
        ws.cell(row, 25).hyperlink = published_url
        ws.cell(row, 25).style = "Hyperlink"
style_data(ws, 4, 3 + len(topics), 1, len(headers), 72)
add_table(ws, f"A3:AD{3 + len(topics)}", "ContentCalendar")
widths = [13, 12, 14, 23, 42, 27, 45, 18, 14, 19, 42, 12, 55, 12, 28, 42, 55, 48, 34, 25, 40, 18, 15, 16, 34, 18, 12, 14, 16, 35]
from openpyxl.utils import get_column_letter
for col, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(col)].width = width
add_list_validation(ws, f"H4:H{3+len(topics)}", "'_lists'!$C$1:$C$6")
add_list_validation(ws, f"I4:I{3+len(topics)}", "'_lists'!$D$1:$D$4")
add_list_validation(ws, f"J4:J{3+len(topics)}", "'_lists'!$G$1:$G$6")
add_list_validation(ws, f"Z4:Z{3+len(topics)}", "'_lists'!$A$1:$A$9")
add_list_validation(ws, f"AA4:AA{3+len(topics)}", "'_lists'!$B$1:$B$4")
ws.conditional_formatting.add(f"L4:L{3+len(topics)}", CellIsRule(operator="greaterThan", formula=[60], fill=PatternFill("solid", fgColor=PALE_RED)))
ws.conditional_formatting.add(f"N4:N{3+len(topics)}", CellIsRule(operator="lessThan", formula=[120], fill=PatternFill("solid", fgColor=PALE_AMBER)))
ws.conditional_formatting.add(f"N4:N{3+len(topics)}", CellIsRule(operator="greaterThan", formula=[165], fill=PatternFill("solid", fgColor=PALE_RED)))
ws.conditional_formatting.add(f"Z4:Z{3+len(topics)}", FormulaRule(formula=['Z4="منتشرشده"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
ws.conditional_formatting.add(f"Z4:Z{3+len(topics)}", FormulaRule(formula=['Z4="در حال نگارش"'], fill=PatternFill("solid", fgColor=PALE_AMBER)))

# Smart brief lookup sheet
ws = wb.create_sheet("بریف مقاله")
setup_sheet(ws, freeze="A4", zoom=95)
title_band(ws, "بریف هوشمند مقاله", "در سلول B3 شناسه محتوا را انتخاب کنید؛ اطلاعات تقویم خودکار نمایش داده می‌شود. بخش‌های پژوهش و نکات تخصصی را پیش از تحویل به نویسنده تکمیل کنید.", 8)
ws["A3"] = "شناسه محتوا"
ws["A3"].font = fa_font(11, True, WHITE); ws["A3"].fill = PatternFill("solid", fgColor=BLUE); ws["A3"].border = border
ws["B3"] = "EC-001"; ws["B3"].fill = PatternFill("solid", fgColor=PALE_CYAN); ws["B3"].font = fa_font(11, True, BLUE); ws["B3"].border = border
add_list_validation(ws, "B3", f"'تقویم محتوا'!$A$4:$A${3+len(topics)}")
brief_fields = [
    ("عنوان مقاله", 5), ("کلمه کلیدی اصلی", 6), ("کلمات فرعی", 7), ("نیت جست‌وجو", 8),
    ("مخاطب اصلی", None), ("مسئله/درد مخاطب", None), ("پاسخ یا وعده اصلی مقاله", None),
    ("SEO Title", 11), ("Meta Description", 13), ("Slug", 15), ("H1", 16), ("ساختار H2", 17),
    ("لینک‌های داخلی", 18), ("منابع ضروری", 19), ("CTA", 20), ("زاویه تمایز با رقبا", None),
    ("مثال واقعی/سناریوی کارگاه", None), ("جدول‌های ضروری", None), ("FAQهای ضروری", None),
    ("ادعاهای نیازمند تأیید", None), ("نکات ویراستار", None),
]
row = 5
for label, calendar_col in brief_fields:
    ws.cell(row, 1, label)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(row, 1).font = fa_font(10, True, INK)
    ws.cell(row, 1).border = border
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    value_cell = ws.cell(row, 2)
    if calendar_col:
        col_letter = get_column_letter(calendar_col)
        value_cell.value = f'=IFERROR(INDEX(\'تقویم محتوا\'!${col_letter}:${col_letter},MATCH($B$3,\'تقویم محتوا\'!$A:$A,0)),"")'
        value_cell.fill = PatternFill("solid", fgColor="F8FAFD")
        value_cell.protection = Protection(locked=True)
    else:
        value_cell.fill = PatternFill("solid", fgColor=PALE_AMBER)
        value_cell.protection = Protection(locked=False)
    value_cell.font = fa_font()
    value_cell.border = border
    value_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    for merged_col in range(3, 9): ws.cell(row, merged_col).border = border
    ws.row_dimensions[row].height = 55 if label not in {"ساختار H2", "نکات ویراستار", "مثال واقعی/سناریوی کارگاه"} else 100
    row += 1
ws.column_dimensions["A"].width = 27
for col in "BCDEFGH": ws.column_dimensions[col].width = 16

# Target pages
ws = wb.create_sheet("صفحات هدف")
setup_sheet(ws, freeze="A4")
headers = ["نوع", "عنوان صفحه", "URL فارسی", "نقش در معماری سئو", "کیوردهای محوری", "CTA اصلی", "لینک از خوشه‌ها", "وضعیت بررسی"]
title_band(ws, "صفحات هدف و Pillar", "این شیت مقصد لینک‌سازی و جلوگیری از تولید مقاله‌های رقیب با صفحات تجاری را مشخص می‌کند.", len(headers))
header_row(ws, 3, headers)
pages = [
    ("Pillar تجاری", "محصول Smart DNC", "/fa/product", "صفحه اصلی تبدیل برای سخت‌افزار و سازگاری", "دستگاه DNC، خرید DNC، Smart DNC", "سفارش/مشاوره", "انتقال فایل، کنترلرها، مقایسه", "بررسی شده"),
    ("Pillar تجاری", "نرم‌افزار ElecIO", "/fa/software", "تبدیل برای مدیریت فایل و گزارش", "نرم افزار DNC، گزارش CNC، مانیتورینگ CNC", "درخواست دمو", "مدیریت برنامه، گزارش، امنیت", "بررسی شده"),
    ("اعتماد", "درباره ما", "/fa/about", "اعتمادسازی برند و تجربه", "ElecIO، شرکت DNC", "آشنایی با تیم", "برند و مطالعات موردی", "بررسی شده"),
    ("تبدیل", "تماس", "/fa/contact", "ثبت درخواست مشاوره و پشتیبانی", "مشاوره CNC، پشتیبانی DNC", "ارسال درخواست", "عیب‌یابی و تصمیم خرید", "بررسی شده"),
    ("Pillar اطلاعاتی", "مجله", "/fa/blog", "هاب تمام خوشه‌های اطلاعاتی", "آموزش CNC، DNC، انتقال فایل", "مطالعه مقالات", "همه خوشه‌ها", "بررسی شده"),
    ("Persona", "مدیران کارگاه", "/fa/for/business-owners", "پیام ارزش اقتصادی و مدیریتی", "کاهش توقف CNC، گزارش تولید", "مشاوره کسب‌وکار", "گزارش و ROI", "بررسی شده"),
    ("Persona", "برنامه‌نویسان CNC", "/fa/for/programmers", "جزئیات انتقال فایل و کنترلر", "انتقال G-code، RS232، DNC", "بررسی سازگاری", "کنترلرها و انتقال", "بررسی شده"),
    ("Persona", "اپراتورها", "/fa/for/operators", "کار روزمره و کاهش خطا", "ارسال برنامه CNC، خطای فایل", "مشاهده روش کار", "آموزش و عیب‌یابی", "بررسی شده"),
]
for r, values in enumerate(pages, 4):
    for c, value in enumerate(values, 1): ws.cell(r, c, value)
    ws.cell(r, 3).hyperlink = values[2]; ws.cell(r, 3).style = "Hyperlink"
style_data(ws, 4, 3 + len(pages), 1, len(headers), 58)
add_table(ws, f"A3:H{3+len(pages)}", "TargetPages")
for col, width in enumerate([18, 28, 32, 42, 45, 24, 34, 18], 1): ws.column_dimensions[get_column_letter(col)].width = width

# Internal linking map
ws = wb.create_sheet("لینک‌سازی داخلی")
setup_sheet(ws, freeze="A4")
headers = ["خوشه مبدأ", "صفحه/خوشه مقصد", "URL مقصد", "انکر پیشنهادی ۱", "انکر پیشنهادی ۲", "جایگاه پیشنهادی", "هدف لینک", "وضعیت"]
title_band(ws, "نقشه لینک‌سازی داخلی", "انکرها را طبیعی و متنوع استفاده کنید؛ هر مقاله باید به Pillar مرتبط و حداقل دو مقاله هم‌خوشه لینک بدهد.", len(headers))
header_row(ws, 3, headers)
links = [
    ("انتقال فایل و DNC", "محصول Smart DNC", "/fa/product", "دستگاه Smart DNC", "راهکار انتقال فایل CNC", "پس از توضیح راهکارها", "تبدیل تجاری"),
    ("مدیریت برنامه", "نرم‌افزار ElecIO", "/fa/software", "نرم‌افزار مدیریت CNC", "مدیریت متمرکز فایل‌ها", "بخش راهکار", "تبدیل تجاری"),
    ("گزارش و بهره‌وری", "نرم‌افزار ElecIO", "/fa/software", "گزارش‌گیری دستگاه CNC", "داشبورد تولید ElecIO", "پس از معرفی شاخص‌ها", "تبدیل تجاری"),
    ("کنترلر FANUC", "راهنمای RS232 فانوک", "/fa/blog/fanuc-rs232-checklist", "تنظیم RS232 فانوک", "چک‌لیست ارتباط فانوک", "بخش تنظیم کنترلر", "عمق موضوعی"),
    ("کنترلر Siemens", "راهنمای زیمنس", "/fa/blog/siemens-program-transfer-workflow", "انتقال برنامه زیمنس", "راهنمای شبکه Siemens", "بخش کنترلرها", "عمق موضوعی"),
    ("امنیت و پشتیبان‌گیری", "راهنمای بکاپ CNC", "/fa/blog/cnc-backup-security-basics", "بکاپ برنامه CNC", "امنیت و بازیابی فایل", "بخش امنیت", "اعتماد/آموزش"),
    ("گزارش و بهره‌وری", "راهنمای توقف CNC", "/fa/blog/cnc-downtime-reporting-guide", "گزارش توقف دستگاه CNC", "تحلیل توقف تولید", "بخش اندازه‌گیری", "عمق موضوعی"),
    ("عیب‌یابی", "تماس", "/fa/contact", "درخواست بررسی فنی", "تماس با تیم ElecIO", "انتهای عیب‌یابی", "تبدیل پشتیبانی"),
    ("اقتصاد تولید", "مدیران کارگاه", "/fa/for/business-owners", "راهکار مدیران کارگاه", "کاهش توقف و خطا", "بخش ROI", "تبدیل Persona"),
    ("آموزش پایه", "راهنمای انتقال فایل", "/fa/blog/cnc-file-transfer-guide", "راهنمای انتقال فایل CNC", "روش ارسال فایل NC", "پس از تعریف مفاهیم", "مسیر مطالعه"),
]
for r, values in enumerate(links, 4):
    for c, value in enumerate((*values, "بررسی نشده"), 1): ws.cell(r, c, value)
    ws.cell(r, 3).hyperlink = values[2]; ws.cell(r, 3).style = "Hyperlink"
style_data(ws, 4, 3 + len(links), 1, len(headers), 55)
add_table(ws, f"A3:H{3+len(links)}", "InternalLinks")
for col, width in enumerate([24, 28, 38, 30, 30, 30, 22, 18], 1): ws.column_dimensions[get_column_letter(col)].width = width
add_list_validation(ws, f"H4:H{3+len(links)}", "'_lists'!$F$1:$F$4")

# Quality checklist
ws = wb.create_sheet("کنترل کیفیت")
setup_sheet(ws, freeze="A4")
headers = ["ردیف", "مرحله", "کنترل ضروری", "معیار قبولی", "وزن", "وضعیت", "مسئول", "یادداشت"]
title_band(ws, "چک‌لیست کنترل کیفیت مقاله", "طبق فایل پرامپت: مقاله باید فارسی، حداقل ۱۵۰۰ کلمه، انسان‌محور، مستند، ساختاریافته و آماده انتشار باشد.", len(headers))
header_row(ws, 3, headers)
checks = [
    ("استراتژی", "نیت جست‌وجو", "مقاله مستقیماً سؤال یا نیاز اصلی کاربر را پاسخ می‌دهد", 5),
    ("استراتژی", "عدم هم‌نوع‌خواری", "با مقاله یا صفحه تجاری موجود رقابت مستقیم ندارد", 5),
    ("تحقیق", "اعتبار منابع", "مشخصات از کاتالوگ/مستند رسمی و ادعاها قابل تأییدند", 5),
    ("تحقیق", "عدم جعل", "هیچ آمار، قیمت، استاندارد یا ادعای رسمی ساخته نشده است", 5),
    ("محتوا", "حداقل طول", "حداقل ۱۵۰۰ واژه فارسی مفید؛ بدون پرگویی", 4),
    ("محتوا", "زبان طبیعی", "روان، حرفه‌ای و متناسب با مخاطب صنعتی ایرانی", 5),
    ("محتوا", "مقدمه", "مسئله، وعده و مسیر مقاله را کوتاه و روشن بیان می‌کند", 3),
    ("ساختار", "یک H1", "فقط یک H1 قوی و شامل مفهوم کیورد اصلی", 5),
    ("ساختار", "H2/H3", "سرفصل‌ها قابل جست‌وجو، منطقی و بدون تکرارند", 4),
    ("ساختار", "جدول مفید", "حداقل یک جدول فقط در صورت افزایش فهم", 3),
    ("ساختار", "مثال عملی", "سناریو یا مثال واقعی از کارگاه CNC دارد", 4),
    ("ساختار", "FAQ", "پرسش‌های واقعی و پاسخ‌های کوتاه و کامل دارد", 4),
    ("ساختار", "جمع‌بندی", "نتیجه‌گیری تصمیم‌ساز و غیرتکراری دارد", 3),
    ("سئو", "استفاده طبیعی کیورد", "کیورد و مترادف‌ها بدون تکرار مصنوعی آمده‌اند", 5),
    ("سئو", "SEO Title", "منحصر‌به‌فرد، واضح و ترجیحاً حداکثر ۶۰ کاراکتر", 4),
    ("سئو", "Meta Description", "منحصر‌به‌فرد، ترغیب‌کننده و حدود ۱۲۰ تا ۱۶۵ کاراکتر", 4),
    ("سئو", "Slug", "کوتاه، انگلیسی، خوانا و پایدار است", 3),
    ("سئو", "لینک داخلی", "به Pillar و دو محتوای مرتبط با انکر طبیعی لینک دارد", 5),
    ("سئو", "لینک خارجی", "فقط منبع معتبر و ضروری؛ بدون زیاده‌روی", 3),
    ("سئو", "دسته و برچسب", "دسته اصلی و تعداد محدود برچسب مرتبط تعیین شده", 2),
    ("تجربه", "خوانایی", "پاراگراف کوتاه، فهرست مناسب و توضیح اصطلاحات دارد", 4),
    ("تجربه", "CTA", "با مرحله قیف و هدف مقاله هماهنگ است", 4),
    ("ویرایش", "نیم‌فاصله و نگارش", "املاء، علائم، اعداد و نیم‌فاصله یکدست‌اند", 3),
    ("ویرایش", "تکرار", "بخش‌های تکراری و جمله‌های رباتیک حذف شده‌اند", 4),
    ("انتشار", "لینک‌ها", "تمام لینک‌ها باز و مقصدشان درست است", 5),
    ("انتشار", "نمایش موبایل", "جدول، تیتر و پاراگراف روی موبایل خواناست", 5),
    ("پس از انتشار", "ایندکس", "Canonical، sitemap و وضعیت index بررسی شده", 4),
    ("پس از انتشار", "پایش", "تاریخ بازبینی GSC و به‌روزرسانی مشخص شده", 3),
]
for i, (stage, check, criterion, weight) in enumerate(checks, 1):
    row = 3 + i
    for col, value in enumerate([i, stage, check, criterion, weight, "بررسی نشده", "", ""], 1): ws.cell(row, col, value)
style_data(ws, 4, 3 + len(checks), 1, len(headers), 48)
add_table(ws, f"A3:H{3+len(checks)}", "QualityChecklist")
for col, width in enumerate([10, 18, 28, 58, 10, 18, 18, 35], 1): ws.column_dimensions[get_column_letter(col)].width = width
add_list_validation(ws, f"F4:F{3+len(checks)}", "'_lists'!$F$1:$F$4")
ws.conditional_formatting.add(f"F4:F{3+len(checks)}", FormulaRule(formula=['F4="تأیید"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
ws.conditional_formatting.add(f"F4:F{3+len(checks)}", FormulaRule(formula=['F4="نیازمند اصلاح"'], fill=PatternFill("solid", fgColor=PALE_RED)))

# Prompt sheet
ws = wb.create_sheet("پرامپت محتوا")
setup_sheet(ws, freeze="A4", zoom=90)
title_band(ws, "پرامپت اختصاصی تولید محتوای ElecIO", "مبتنی بر Blog Prompt.docx و متناسب با حوزه CNC/DNC. مقادیر داخل {{ }} را از بریف مقاله جایگزین کنید.", 8)
prompt_sections = [
    ("نقش", "تو یک نویسنده متخصص سئوی فارسی، استراتژیست محتوای B2B صنعتی و آشنا با CNC، DNC، کنترلرهای FANUC، Siemens، Heidenhain و ارتباط RS232/شبکه هستی."),
    ("ورودی‌ها", "کلمه کلیدی اصلی: {{کلمه کلیدی اصلی}}\nکلمات فرعی: {{کلمات فرعی}}\nعنوان: {{عنوان}}\nنیت جست‌وجو: {{نیت}}\nمخاطب: {{مخاطب}}\nCTA: {{CTA}}\nمنابع تأییدشده: {{منابع}}"),
    ("هدف", "یک مقاله فارسی کامل، کاربردی و قابل انتشار برای وب‌سایت ElecIO بنویس که مسئله واقعی مدیر کارگاه، برنامه‌نویس یا اپراتور CNC را حل کند. انسان‌محور بنویس و از Keyword Stuffing پرهیز کن."),
    ("محدودیت صحت", "هیچ آمار، قیمت، قانون، استاندارد، سازگاری کنترلر یا مشخصه فنی را اختراع نکن. برای مشخصات ElecIO فقط از کاتالوگ رسمی، صفحات محصول و اطلاعات تأییدشده استفاده کن. اگر داده‌ای موجود نیست، شفاف بگو نیازمند تأیید فنی است."),
    ("طول و زبان", "حداقل ۱۵۰۰ واژه مفید و کاملاً فارسی. اصطلاحات انگلیسی تخصصی را در صورت نیاز کنار معادل فارسی بیاور. از لحن رباتیک، مقدمه طولانی و تکرار مصنوعی خودداری کن."),
    ("ساختار", "یک H1 قوی، مقدمه کوتاه، H2های قابل جست‌وجو، H3 برای مراحل/مثال‌ها/اشتباهات، حداقل یک جدول مفید در صورت نیاز، مثال عملی کارگاهی، نکات اجرایی، FAQ واقعی و جمع‌بندی تصمیم‌ساز."),
    ("SEO Metadata", "در ابتدای خروجی این موارد را بده: SEO Title، Meta Description، Slug انگلیسی پیشنهادی، Focus Keyword، Secondary Keywords، دسته پیشنهادی، برچسب‌های محدود و مرتبط."),
    ("لینک‌سازی", "لینک داخلی را با قالب [پیشنهاد لینک داخلی: انکر — URL] و فقط در جای طبیعی پیشنهاد بده. حداقل یک لینک به Pillar مرتبط و دو لینک به مقالات هم‌خوشه. لینک خارجی فقط به منابع رسمی و معتبر، در صورت ارزش واقعی."),
    ("ممنوعیت تصویر", "تصویر نساز، Placeholder تصویر، محل تصویر، کپشن یا Alt Text پیشنهاد نده."),
    ("خروجی Word", "در صورت درخواست فایل Word: بدون Header و Footer، با Heading Style صحیح، فاصله خطوط ۲.۲، پاراگراف‌های خوانا، فهرست‌های استاندارد و جدول‌های حرفه‌ای تحویل بده."),
    ("کنترل نهایی", "پیش از تحویل بررسی کن: حداقل ۱۵۰۰ واژه، پاسخ کامل به نیت، ساختار درست، استفاده طبیعی کیورد، منابع قابل اتکا، FAQ، جمع‌بندی، لینک داخلی، CTA و نبود ادعای ساختگی."),
    ("قالب خروجی", "۱) متادیتای SEO\n۲) H1\n۳) مقدمه\n۴) بدنه با H2/H3\n۵) جدول/چک‌لیست مفید\n۶) FAQ\n۷) جمع‌بندی و CTA\n۸) لینک‌های داخلی و خارجی پیشنهادی"),
]
row = 4
for section, content in prompt_sections:
    ws.cell(row, 1, section)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
    ws.cell(row, 1).font = fa_font(11, True, WHITE)
    ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row, 1).border = border
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    ws.cell(row, 2, content)
    ws.cell(row, 2).fill = PatternFill("solid", fgColor="F8FAFD")
    ws.cell(row, 2).font = fa_font(10)
    ws.cell(row, 2).alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    ws.cell(row, 2).border = border
    for col in range(3, 9): ws.cell(row, col).border = border
    ws.row_dimensions[row].height = 70 if len(content) < 250 else 105
    row += 1
ws.column_dimensions["A"].width = 22
for col in "BCDEFGH": ws.column_dimensions[col].width = 17

# Guide
ws = wb.create_sheet("راهنما", 0)
setup_sheet(ws, freeze="A4", zoom=95)
title_band(ws, "راهنمای استفاده از فایل", "این فایل برای برنامه‌ریزی، نگارش، انتشار و بهینه‌سازی مستمر محتوای فارسی ElecIO طراحی شده است.", 8)
guide = [
    ("۱. داده واقعی وارد کنید", "در «نقشه کیورد»، حجم جست‌وجو، سختی، رتبه، کلیک و ایمپرشن را فقط از Google Search Console، Keyword Planner یا ابزار معتبر وارد کنید."),
    ("۲. اولویت را تعیین کنید", "ابتدا موضوعات با نیت تجاری، ارتباط مستقیم با محصول و شکاف محتوایی واقعی را انتخاب کنید. داشبورد وضعیت را خلاصه می‌کند."),
    ("۳. بریف بسازید", "در «بریف مقاله» شناسه‌ای مثل EC-006 را انتخاب کنید؛ عنوان، کیورد، متا، ساختار و لینک‌ها خودکار نمایش داده می‌شوند. خانه‌های زرد را تیم فنی/سئو تکمیل کند."),
    ("۴. پرامپت را استفاده کنید", "پرامپت اختصاصی را با داده‌های بریف کامل کنید. در موضوعات فنی، کاتالوگ و تأیید تیم محصول باید بر متن AI مقدم باشد."),
    ("۵. کنترل کیفیت", "قبل از انتشار تمام موارد وزن ۵ در چک‌لیست باید تأیید شوند. سپس نگارش، موبایل، لینک‌ها و متادیتا بررسی شود."),
    ("۶. لینک‌سازی", "هر مقاله به صفحه Pillar مرتبط و دو مقاله هم‌خوشه لینک بدهد. انکر را طبیعی و متنوع انتخاب کنید."),
    ("۷. پس از انتشار", "لینک انتشار، وضعیت، درصد پیشرفت و تاریخ به‌روزرسانی را ثبت کنید؛ پس از ۲۸ تا ۴۲ روز داده GSC را بررسی و محتوا را بهبود دهید."),
    ("قاعده مهم", "این فایل پیشنهاد استراتژیک است، نه داده تحقیق کیورد. حجم و رتبه تخمینی ساخته نشده و باید با داده واقعی بازار ایران تکمیل شود."),
]
row = 4
for heading, body in guide:
    ws.cell(row, 1, heading)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(row, 1).font = fa_font(11, True, BLUE)
    ws.cell(row, 1).border = border
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    ws.cell(row, 2, body)
    ws.cell(row, 2).font = fa_font(10)
    ws.cell(row, 2).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.cell(row, 2).border = border
    for col in range(3, 9): ws.cell(row, col).border = border
    ws.row_dimensions[row].height = 58
    row += 1
ws.column_dimensions["A"].width = 25
for col in "BCDEFGH": ws.column_dimensions[col].width = 17

# Workbook-level details
wb.properties.title = "ElecIO SEO Content Planner"
wb.properties.subject = "Persian SEO keyword map, content calendar, brief and quality workflow"
wb.properties.creator = "ElecIO"
wb.properties.keywords = "ElecIO, SEO, CNC, DNC, Persian content"
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)
print(f"Created {OUT} with {len(topics)} planned topics and {len(wb.sheetnames)-1} visible sheets.")
