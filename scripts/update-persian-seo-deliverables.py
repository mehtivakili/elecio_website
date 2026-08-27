from __future__ import annotations

from copy import copy
from datetime import date, timedelta
from html import escape
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
BOOK = OUT / 'Elecio_SEO_Content_Planner_FA.xlsx'
REVISED_BOOK = OUT / 'Elecio_SEO_Content_Planner_FA_Revised.xlsx'
HTML = OUT / 'Elecio_SEO_Persian_URL_and_Image_Plan.html'
QUEUE = OUT / 'Persian_SEO_Publication_Queue'
BASE_PROMPT = 'تصویر افقی جنرال و حرفه‌ای برای مقاله صنعتی، واقع‌گرایانه، نور تمیز، ترکیب‌بندی خلوت، بدون نوشته، لوگو و واترمارک، نسبت ۳:۲، خروجی دقیق ۶۰۰×۴۰۰ پیکسل'

# keyword, cluster, related, intent, funnel, format, pillar, priority, slug, title, meta, h2, image title, alt, scene
DATA = [
('جی کد','برنامه‌نویسی CNC','G-code، کدهای G و M، برنامه CNC','اطلاعاتی','آگاهی','راهنمای جامع','/fa/for/programmers','خیلی بالا','جی-کد-چیست','جی کد چیست؟ راهنمای کامل G-code برای دستگاه CNC','جی کد زبان فرمان دستگاه CNC است. در این راهنمای کاربردی با ساختار G-code، کدهای اصلی، نمونه برنامه، خطاهای رایج و روش انتقال امن فایل آشنا شوید.','جی کد چیست؟ | ساختار یک برنامه G-code | کدهای G و M پرکاربرد | نمونه عملی | خطاهای رایج | پرسش‌های متداول','ارتباط جی کد با کنترلر و دستگاه CNC','نمای کنترلر CNC، مسیر ابزار و قطعه فلزی ماشین‌کاری‌شده','کنترلر CNC در کنار نمایش گرافیکی مسیر ابزار روی قطعه فلزی و دستگاه CNC در پس‌زمینه'),
('cnc دستگاه','مبانی CNC','دستگاه CNC، ماشین CNC، انواع CNC','اطلاعاتی/تجاری','بررسی','راهنمای جامع','/fa/product','خیلی بالا','دستگاه-cnc-چیست','دستگاه CNC چیست؟ انواع، کاربرد و نکات انتخاب','دستگاه CNC چگونه کار می‌کند و چه انواعی دارد؟ اجزا، کاربردها، تفاوت مدل‌ها، هزینه‌های پنهان و معیارهای انتخاب دستگاه مناسب کارگاه را بررسی می‌کنیم.','دستگاه CNC چیست؟ | اجزای اصلی | انواع دستگاه CNC | کاربردها | معیارهای انتخاب | نگهداری و ایمنی','نمای عمومی دستگاه CNC در محیط صنعتی','دستگاه CNC مدرن در یک کارگاه صنعتی تمیز','یک دستگاه CNC مدرن و کامل در کارگاه صنعتی تمیز، قطعه فلزی و ابزار برش قابل مشاهده'),
('فرز CNC','فرز CNC','دستگاه فرز CNC، فرز سه محور، فرز صنعتی','اطلاعاتی/تجاری','بررسی','راهنمای فنی','/fa/product','خیلی بالا','فرز-cnc-چیست','فرز CNC چیست؟ راهنمای انواع، اجزا و کاربردها','با فرز CNC، انواع سه تا پنج محور، اجزای اصلی، کاربردهای صنعتی و تفاوت آن با تراش CNC آشنا شوید و معیارهای انتخاب مدل مناسب را ببینید.','فرز CNC چگونه کار می‌کند؟ | انواع محور | اجزای اصلی | کاربردها | تفاوت فرز و تراش | معیارهای انتخاب','فرز CNC هنگام ماشین‌کاری قطعه فلزی','اسپیندل فرز CNC در حال براده‌برداری از قطعه فلزی','نمای نزدیک اسپیندل فرز CNC هنگام براده‌برداری دقیق از قطعه آلومینیومی، پاشش کنترل‌شده خنک‌کننده'),
('اتوماسیون صنعتی','اتوماسیون صنعتی','کنترل صنعتی، هوشمندسازی کارخانه، مانیتورینگ تولید','اطلاعاتی/تجاری','بررسی','راهنمای جامع','/fa/software','بالا','اتوماسیون-صنعتی-چیست','اتوماسیون صنعتی چیست؟ اجزا، سطوح و کاربردها','اتوماسیون صنعتی از حسگر و PLC تا مانیتورینگ و نرم‌افزار تولید چگونه کار می‌کند؟ اجزا، سطوح، مزایا، ریسک‌ها و مسیر شروع در کارگاه را بخوانید.','تعریف اتوماسیون صنعتی | اجزای اصلی | سطوح اتوماسیون | PLC و شبکه | مزایا و ریسک‌ها | نقشه شروع','زنجیره اتوماسیون در کارخانه هوشمند','خط تولید خودکار با ربات، حسگر و تابلو کنترل صنعتی','خط تولید خودکار در کارخانه هوشمند شامل ربات صنعتی، حسگرها، نوار نقاله و تابلو کنترل، بدون رابط کاربری خوانا'),
('plc','کنترل صنعتی','PLC چیست، کنترل‌گر منطقی برنامه‌پذیر، برنامه نویسی PLC','اطلاعاتی','آگاهی','آموزش فنی','/fa/software','بالا','plc-چیست','PLC چیست؟ راهنمای کنترل‌گر منطقی برنامه‌پذیر','PLC قلب بسیاری از سامانه‌های کنترل صنعتی است. ساختار، ورودی و خروجی، چرخه اسکن، زبان‌های برنامه‌نویسی و معیارهای انتخاب PLC را ساده و فنی بیاموزید.','PLC چیست؟ | اجزای PLC | چرخه اسکن | ورودی و خروجی | زبان‌های برنامه‌نویسی | انتخاب PLC','PLC و اجزای تابلو کنترل صنعتی','ماژول PLC، ورودی و خروجی و سیم‌کشی منظم تابلو برق','نمای حرفه‌ای یک ماژول PLC در تابلو کنترل صنعتی با ماژول‌های ورودی و خروجی و سیم‌کشی منظم'),
('تراش cnc','تراش CNC','دستگاه تراش CNC، تراشکاری CNC، محور C','اطلاعاتی/تجاری','بررسی','راهنمای فنی','/fa/product','خیلی بالا','تراش-cnc-چیست','تراش CNC چیست؟ اجزا، عملیات و نکات انتخاب','تراش CNC چگونه قطعات دقیق تولید می‌کند؟ اجزا، محور‌ها، عملیات رایج، تفاوت با فرز CNC، ابزارها و معیارهای انتخاب دستگاه مناسب را بررسی کنید.','تراش CNC چگونه کار می‌کند؟ | اجزای دستگاه | عملیات تراشکاری | ابزار و آفست | تفاوت با فرز | معیارهای انتخاب','عملیات تراشکاری با دستگاه تراش CNC','سه‌نظام و ابزار دستگاه تراش CNC در حال تولید قطعه فلزی','نمای نزدیک سه‌نظام و ابزار دستگاه تراش CNC هنگام شکل‌دهی قطعه فلزی استوانه‌ای با براده‌های کنترل‌شده'),
('ماشین‌کاری','ماشین‌کاری','ماشین کاری چیست، براده برداری، فرایندهای ساخت','اطلاعاتی','آگاهی','راهنمای جامع','/fa/blog','بالا','ماشین-کاری-چیست','ماشین‌کاری چیست؟ فرایندها، ابزارها و کاربردها','ماشین‌کاری مجموعه فرایندهای ساخت با حذف ماده است. تراش، فرز، سوراخ‌کاری، سنگ‌زنی، پارامترهای برش و نقش CNC در تولید دقیق را مرور می‌کنیم.','تعریف ماشین‌کاری | فرایندهای اصلی | ابزار برش | پارامترهای برش | دقت و کیفیت سطح | نقش CNC','نمای عمومی فرایندهای ماشین‌کاری فلز','ابزارهای فرز، تراش و سوراخ‌کاری در کنار قطعات ماشین‌کاری‌شده','چیدمان صنعتی منظم از ابزار فرز، ابزار تراش، مته و چند قطعه فلزی دقیق ماشین‌کاری‌شده روی میز کار'),
('مسترکم','نرم‌افزار CAM','Mastercam، آموزش مسترکم، برنامه نویسی CNC','اطلاعاتی/تجاری','بررسی','راهنمای نرم‌افزار','/fa/for/programmers','بالا','مسترکم-چیست','مسترکم چیست؟ کاربرد Mastercam در برنامه‌نویسی CNC','مسترکم چگونه مسیر ابزار و جی کد تولید می‌کند؟ محیط‌ها، گردش کار CAD/CAM، پست‌پروسسور، شبیه‌سازی و نکات انتخاب و یادگیری Mastercam را بخوانید.','مسترکم چیست؟ | گردش کار CAD/CAM | مسیر ابزار | پست‌پروسسور | شبیه‌سازی | نکات یادگیری','طراحی مسیر ابزار CAM برای دستگاه CNC','مدل سه‌بعدی قطعه و مسیر ابزار رنگی روی نمایشگر مهندسی','ایستگاه کاری مهندسی با مدل سه‌بعدی قطعه و مسیر ابزار CAM رنگی روی نمایشگر، دستگاه CNC محو در پس‌زمینه، بدون برند و متن خوانا'),
('پاورمیل','نرم‌افزار CAM','PowerMill، آموزش پاورمیل، ماشین کاری سه بعدی','اطلاعاتی/تجاری','بررسی','راهنمای نرم‌افزار','/fa/for/programmers','بالا','پاورمیل-چیست','پاورمیل چیست؟ کاربرد PowerMill در ماشین‌کاری CNC','PowerMill برای ساخت مسیر ابزار قطعات پیچیده چگونه استفاده می‌شود؟ استراتژی‌های خشن‌کاری و پرداخت، شبیه‌سازی، پست‌پروسسور و گردش کار را بررسی کنید.','پاورمیل چیست؟ | آماده‌سازی مدل | خشن‌کاری | پرداخت | شبیه‌سازی | خروجی G-code','شبیه‌سازی مسیر ابزار قطعه پیچیده CNC','مسیر ابزار سه‌بعدی خشن‌کاری و پرداخت روی یک قطعه صنعتی','مدل سه‌بعدی قطعه صنعتی پیچیده با مسیرهای ابزار خشن‌کاری و پرداخت رنگی در محیط CAM، بدون برند و متن خوانا'),
('خرید فرز','خرید دستگاه CNC','خرید فرز CNC، قیمت فرز، فرز نو یا کارکرده','خرید','تصمیم','راهنمای خرید','/fa/contact','خیلی بالا','راهنمای-خرید-فرز-cnc','راهنمای خرید فرز CNC؛ چک‌لیست فنی و هزینه‌ها','پیش از خرید فرز CNC چه چیزهایی را بررسی کنیم؟ نوع کنترلر، محورها، اسپیندل، دقت، خدمات، قطعات، تست عملی و هزینه کل مالکیت را مقایسه کنید.','نیاز کارگاه | مشخصات فنی | کنترلر و انتقال فایل | تست سلامت | نو یا کارکرده | هزینه کل مالکیت | چک‌لیست خرید','بررسی فنی دستگاه پیش از خرید فرز CNC','کارشناس در حال بررسی فنی دستگاه فرز CNC در کارگاه','کارشناس صنعتی در حال بررسی اسپیندل و کنترلر یک دستگاه فرز CNC در کارگاه، فضای مقایسه و ارزیابی خرید'),
]

KEYS = ['keyword','cluster','related','intent','funnel','format','pillar','priority','slug','title','meta','h2','image_title','image_alt','scene']
ITEMS = [dict(zip(KEYS, row)) for row in DATA]
for i, x in enumerate(ITEMS, 1):
    x.update(id=f'EC-{i:03}', url=f"https://elecio.co/fa/blog/{x['slug']}", image_file=f"{x['slug']}-تصویر-شاخص.webp", prompt=f"{x['scene']}؛ {BASE_PROMPT}", folder=f"EC-{i:03}-{x['slug']}")

def clone(source, target):
    if source.has_style: target._style = copy(source._style)
    target.number_format, target.alignment = source.number_format, copy(source.alignment)

def clone_row(ws, source, target):
    for col in range(1, ws.max_column + 1): clone(ws.cell(source, col), ws.cell(target, col))
    ws.row_dimensions[target].height = ws.row_dimensions[source].height

def clear(ws, start=4):
    for row in ws.iter_rows(min_row=start):
        for cell in row: cell.value = None

def operation_sheet(wb):
    if 'مسیر عملیاتی' in wb.sheetnames: del wb['مسیر عملیاتی']
    ws = wb.create_sheet('مسیر عملیاتی', 2); ws.sheet_view.rightToLeft = True
    headers = ['مرحله','مسئول','ورودی','اقدام دقیق','معیار پذیرش','محل ثبت']
    steps = [
    ('۱. انتخاب موضوع','مدیر محتوا','نقشه کیورد','موضوع با نیت روشن انتخاب و هم‌پوشانی با صفحات قبلی بررسی شود.','یک ID و keyword اصلی','نقشه کیورد'),
    ('۲. تثبیت URL','سئو','عنوان و keyword','هسته موضوع حفظ و واژه‌های کم‌ارزش حذف؛ کلمات با خط فاصله جدا شوند.','URL یکتا، کوتاه و بدون تاریخ','تقویم محتوا'),
    ('۳. نگارش Word','نویسنده','بریف پوشه موضوع','مقدمه پاسخ‌محور، H2ها، مثال، منبع، FAQ و CTA نوشته شود.','متن کامل و بازبینی‌شده','Word همان پوشه'),
    ('۴. تولید تصویر','طراح/AI','پرامپت اکسل','تصویر بدون متن/لوگو ساخته، ۶۰۰×۴۰۰ و WebP ذخیره شود.','تصویر، نام و alt صحیح','پوشه images'),
    ('۵. ورود به CMS','ناشر','Word و تصویر','Title، Meta، slug، بدنه، تصویر و alt وارد و preview بررسی شود.','کنترل‌های SEO کامل','پنل نویسنده'),
    ('۶. صف انتشار','مدیر محتوا','Draft آماده','Active Queue روشن و فاصله ۲، ۳ یا ۴ روز انتخاب شود.','تاریخ آینده ثبت شده','صف انتشار'),
    ('۷. پس از انتشار','سئو','URL نهایی','۲۰۰، canonical، index، sitemap، تصویر و لینک داخلی بررسی شود.','بدون redirect و قابل ایندکس','کنترل کیفیت'),
    ('۸. بازبینی','سئو','داده GSC','پس از ۲۸ تا ۴۲ روز query، CTR و position ثبت و تصمیم‌گیری شود.','تصمیم بهبود مستند','داشبورد')]
    thin = Side(style='thin', color='CCD5E2')
    for c, v in enumerate(headers, 1):
        z=ws.cell(1,c,v); z.fill=PatternFill('solid',fgColor='172033'); z.font=Font(color='FFFFFF',bold=True); z.alignment=Alignment(horizontal='center')
    for r, vals in enumerate(steps, 2):
        for c,v in enumerate(vals,1):
            z=ws.cell(r,c,v); z.alignment=Alignment(vertical='top',wrap_text=True); z.border=Border(left=thin,right=thin,top=thin,bottom=thin)
            if r%2==0: z.fill=PatternFill('solid',fgColor='EEF4FF')
    for c,w in enumerate([18,14,23,60,42,24],1): ws.column_dimensions[ws.cell(1,c).column_letter].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref='A1:F9'

def workbook():
    wb=load_workbook(BOOK); source='فهرست منتخب کاربر از جست‌وجوی دستی Google — seo_keywords.txt — بدون ادعای حجم یا سختی'
    ws=wb['نقشه کیورد']; clear(ws)
    for i,x in enumerate(ITEMS,1):
        r=i+3; clone_row(ws,4,r); vals=[f'KW-{i:03}',x['keyword'],x['cluster'],x['related'],x['intent'],x['funnel'],x['format'],x['pillar'],x['priority'],source,None,None,None,10,None,None,None,'نامشخص',x['id'],date(2026,8,28),'Volume، KD و رتبه فقط با داده واقعی GSC/Keyword Planner تکمیل شود.']
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
    ws=wb['تقویم محتوا']; headers=['URL نهایی فارسی','عنوان تصویر شاخص','پرامپت خلاصه تولید تصویر','ابعاد/فرمت','نام فایل تصویر','متن جایگزین (alt)','مسیر پوشه تحویل']
    for c,h in enumerate(headers,31): clone(ws.cell(3,30),ws.cell(3,c)); ws.cell(3,c).value=h; clone(ws.cell(4,30),ws.cell(4,c)); ws.column_dimensions[ws.cell(3,c).column_letter].width=65 if c==33 else 28
    clear(ws); start=date(2026,9,1); days=['دوشنبه','سه‌شنبه','چهارشنبه','پنج‌شنبه','جمعه','شنبه','یکشنبه']
    for i,x in enumerate(ITEMS,1):
        r=i+3; clone_row(ws,4,r); pub=start+timedelta(days=(i-1)*3); internal=f"{x['pillar']} | /fa/blog | "+(f"/fa/blog/{ITEMS[i-2]['slug']}" if i>1 else '/fa/product')
        vals=[x['id'],(i+1)//2,days[pub.weekday()],x['cluster'],x['title'],x['keyword'],x['related'],x['intent'],x['funnel'],x['format'],x['title'],f'=LEN(K{r})',x['meta'],f'=LEN(M{r})',x['slug'],x['title'],x['h2'],internal,'منابع رسمی و تجربه فنی قابل‌تأیید','مطالعه مقاله مرتبط یا درخواست مشاوره',f"پاسخ کامل به نیت «{x['keyword']}» بدون کیورد استافینگ",None,None,pub,x['url'],'آماده نگارش',x['priority'],70,date(2026,8,28),'۱ شاخص ۶۰۰×۴۰۰ + تصویر داخلی فقط در صورت افزایش فهم',x['url'],x['image_title'],x['prompt'],'600×400 / WebP / 3:2',x['image_file'],x['image_alt'],f"deliverables/Persian_SEO_Publication_Queue/{x['folder']}"]
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
    ws.auto_filter.ref=f'A3:AK{len(ITEMS)+3}'
    ws=wb['لینک‌سازی داخلی']; clear(ws)
    for i,x in enumerate(ITEMS,1):
        r=i+3; clone_row(ws,4,r); y=ITEMS[i%len(ITEMS)]; vals=[x['cluster'],f"مقاله {y['keyword']}",f"/fa/blog/{y['slug']}",y['keyword'],f"راهنمای {y['keyword']}",'پس از تعریف مفهوم مرتبط','عمق موضوعی','بررسی نشده']
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
    ws=wb['کنترل کیفیت']; additions=[('تصویر','تصویر شاخص ۶۰۰×۴۰۰','جنرال، مرتبط، بدون نوشته/لوگو، ۳:۲ و WebP',5),('تصویر','نام فایل و alt','نام توصیفی؛ alt توصیف واقعی تصویر و بدون تکرار مصنوعی keyword',5),('سئو','URL فارسی نهایی','کوتاه، پایدار، معنادار، بدون تاریخ و پارامتر',5),('انتشار','صف خودکار','متن، keyword، Title، Meta، دسته، تصویر، alt و index کامل',5)]
    existing={str(ws.cell(r,3).value):r for r in range(4,ws.max_row+1)}
    for stage,control,criterion,weight in additions:
        r=existing.get(control) or ws.max_row+1; clone_row(ws,4,r)
        for c,v in enumerate([r-3,stage,control,criterion,weight,'بررسی نشده',None,None],1): ws.cell(r,c).value=v
    wb['راهنما']['A11']='منبع کیوردها و قاعده URL'; wb['راهنما']['B11']='۱۰ seed فایل seo_keywords.txt وارد شده‌اند. URL فارسی کوتاه و پایدار است؛ اجزای اصلی با خط فاصله می‌آیند و کلمات کم‌ارزش حذف می‌شوند. Volume/KD فقط با داده واقعی تکمیل شود.'
    operation_sheet(wb); wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True
    try: wb.save(BOOK)
    except PermissionError: wb.save(REVISED_BOOK)

def p(text,style=None):
    s=f'<w:pStyle w:val="{style}"/>' if style else ''
    return f'<w:p><w:pPr>{s}<w:bidi/></w:pPr><w:r><w:rPr><w:rtl/></w:rPr><w:t xml:space="preserve">{escape(text,quote=False)}</w:t></w:r></w:p>'

def docx(x):
    folder=QUEUE/x['folder']; images=folder/'images'; images.mkdir(parents=True,exist_ok=True)
    body=[p(x['title'],'Title'),p(f"کلمه کلیدی: {x['keyword']}"),p(f"URL نهایی: {x['url']}"),p('مشخصات سئو','Heading1'),p(f"عنوان سئو: {x['title']}"),p(f"Meta: {x['meta']}"),p(f"کلمات مرتبط: {x['related']}"),p('تصویر شاخص','Heading1'),p(f"عنوان: {x['image_title']}"),p(f"نام فایل: {x['image_file']}"),p('ابعاد: ۶۰۰×۴۰۰، WebP، نسبت ۳:۲'),p(f"alt: {x['image_alt']}"),p(f"پرامپت: {x['prompt']}"),p('متن مقاله','Heading1'),p('مقدمه: در ۸۰ تا ۱۲۰ کلمه پاسخ مستقیم، مسئله کاربر و وعده دقیق مقاله را بیان کنید.')]
    for h in x['h2'].split(' | '): body += [p(h,'Heading2'),p(f"[این بخش با تعریف روشن، مثال صنعتی و تجربه یا منبع قابل‌تأیید تکمیل شود. از تکرار مصنوعی «{x['keyword']}» خودداری کنید.]")]
    body += [p('لینک‌ها و CTA','Heading1'),p(f"صفحه اصلی مرتبط: {x['pillar']} | هاب: /fa/blog | CTA: مطالعه مقاله مرتبط یا درخواست مشاوره"),p('کنترل نهایی','Heading1'),p('□ صحت فنی و منابع  □ Title و Meta  □ URL و canonical  □ تصویر و alt  □ لینک داخلی  □ FAQ  □ موبایل  □ صف انتشار')]
    document='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'+''.join(body)+'<w:sectPr/></w:body></w:document>'
    types='<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>'
    rels='<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
    styles='<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:sz w:val="22"/><w:rtl/></w:rPr></w:rPrDefault></w:docDefaults><w:style w:type="paragraph" w:styleId="Title"><w:rPr><w:b/><w:sz w:val="36"/><w:color w:val="1464F4"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Heading1"><w:rPr><w:b/><w:sz w:val="28"/></w:rPr></w:style><w:style w:type="paragraph" w:styleId="Heading2"><w:rPr><w:b/><w:sz w:val="24"/></w:rPr></w:style></w:styles>'
    with ZipFile(folder/f"{x['id']}-متن-و-بریف-مقاله.docx",'w',ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml',types); z.writestr('_rels/.rels',rels); z.writestr('word/document.xml',document); z.writestr('word/styles.xml',styles)
    (images/'راهنمای-تصویر.txt').write_text(f"عنوان: {x['image_title']}\nنام فایل: {x['image_file']}\nابعاد: 600×400 px\nفرمت: WebP\nalt: {x['image_alt']}\n\nپرامپت:\n{x['prompt']}\n",encoding='utf-8')

def folders():
    QUEUE.mkdir(parents=True,exist_ok=True)
    for x in ITEMS: docx(x)

def html():
    rows=''.join(f"<tr><td>{x['id']}</td><td>{escape(x['keyword'])}</td><td>{escape(x['title'])}</td><td class=ltr>{escape(x['url'])}</td><td>{escape(x['image_title'])}<br><small>{escape(x['image_file'])}</small></td><td>{escape(x['prompt'])}</td></tr>" for x in ITEMS)
    HTML.write_text(f'''<!doctype html><html lang=fa dir=rtl><head><meta charset=utf-8><title>راهنمای عملیاتی سئوی فارسی ElecIO</title><style>@font-face{{font-family:v;src:url('../public/fonts/vazirmatn-400.woff2')}}@font-face{{font-family:v;src:url('../public/fonts/vazirmatn-700.woff2');font-weight:700}}@page{{size:A4;margin:14mm}}*{{box-sizing:border-box}}body{{font-family:v,Arial;color:#172033;line-height:1.8;font-size:11px}}h1{{color:#1464f4;margin:0}}h2{{border-bottom:2px solid #e6edf7;padding-bottom:3px}}.lead,small{{color:#536078}}.note{{background:#eef4ff;border-right:4px solid #1464f4;padding:10px 14px;border-radius:8px}}.warn{{background:#fff6df;border-color:#f4aa16}}table{{width:100%;border-collapse:collapse;font-size:8px}}tr{{page-break-inside:avoid}}th,td{{border:1px solid #ccd5e2;padding:5px;vertical-align:top}}th{{background:#172033;color:white}}.ltr{{direction:ltr;text-align:left;overflow-wrap:anywhere}}.page{{page-break-before:always}}code,pre{{direction:ltr}}a{{color:#1464f4}}</style></head><body>
<h1>راهنمای کامل URL، تصویر و انتشار محتوای فارسی ElecIO</h1><p class=lead>نسخه ۱۴۰۵/۰۶/۰۶ — دامنه مرجع: https://elecio.co — ویژه فارسی</p><div class=note><b>تصمیم:</b> URL مقاله فارسی با اسلاگ فارسی کوتاه ساخته می‌شود. در عنوان بلند فقط هسته موضوع و عبارت متمایزکننده حفظ و واژه‌های کم‌ارزش حذف می‌شوند.</div>
<h2>۱. قاعده URL فارسی</h2><ul><li>الگو: <code>https://elecio.co/fa/blog/کلمه-اصلی-موضوع</code></li><li>بین واژه‌ها فقط خط فاصله؛ بدون فاصله، underscore، نیم‌فاصله، علامت سؤال، تاریخ یا پارامتر.</li><li>URL یکتا و پس از انتشار ثابت باشد. keyword را طبیعی نگه دارید و همه کلمات عنوان را وارد نکنید.</li><li><code>%D8%...</code> فقط percent-encoding همان حروف فارسی است، نه URL دیگر.</li><li>canonical به URL نهایی با پاسخ ۲۰۰ اشاره کند. اگر URL تغییر کرد، redirect دائمی مستقیم، sitemap و لینک داخلی اصلاح شوند.</li></ul><div class='note warn'><b>انتظار درست:</b> فارسی بودن slug تضمین رتبه نیست؛ کیفیت پاسخ، نیت جست‌وجو، لینک داخلی، indexability، اعتبار و تجربه صفحه مهم‌ترند.</div>
<h2>۲. مسیر عملیاتی</h2><ol><li>ردیف و ID را از اکسل انتخاب و keyword، عنوان و URL را قفل کنید.</li><li>متن را در Word همان پوشه بنویسید و تصاویر را فقط در images بگذارید.</li><li>مقدمه پاسخ‌محور، H2، مثال فنی، منبع، FAQ، لینک داخلی و CTA را کامل کنید.</li><li>تصویر را از پرامپت ردیف بسازید؛ ۶۰۰×۴۰۰، WebP، بدون متن/لوگو و با نام تعیین‌شده.</li><li>در CMS فیلدها و پیش‌نمایش را کنترل و Draft را وارد Active Queue دو/سه/چهارروزه کنید.</li><li>پس از انتشار ۲۰۰، canonical، index، sitemap و لینک‌ها را بررسی و پس از ۲۸–۴۲ روز GSC را ثبت کنید.</li></ol>
<h2>۳. پوشه‌بندی و Word</h2><pre>Persian_SEO_Publication_Queue/
  EC-001-جی-کد-چیست/
    EC-001-متن-و-بریف-مقاله.docx
    images/
      راهنمای-تصویر.txt
      جی-کد-چیست-تصویر-شاخص.webp</pre><p>Word بریف سئو، ساختار، تصویر و چک‌لیست دارد. متن نهایی باید جای placeholderها را در همان فایل بگیرد.</p>
<h2>۴. استاندارد تصویر</h2><ul><li>شاخص دقیقاً ۶۰۰×۴۰۰، ۳:۲، WebP و ترجیحاً کمتر از ۱۵۰KB بدون افت محسوس.</li><li>یک تصویر جنرال؛ ۲ تا ۴ تصویر داخلی فقط اگر فهم فنی را بیشتر می‌کند.</li><li>alt مشاهدات واقعی تصویر را توصیف کند. تصویر AI از نظر ابزار، دستگاه، دست، نوشته و لوگوی جعلی بازبینی شود.</li></ul>
<div class=page><h2>۵. نقشه موضوع‌ها، URL و تصویر</h2><table><thead><tr><th>ID</th><th>keyword</th><th>عنوان</th><th>URL</th><th>عنوان/نام تصویر</th><th>پرامپت ۶۰۰×۴۰۰</th></tr></thead><tbody>{rows}</tbody></table><h2>۶. چک‌لیست انتشار</h2><ol><li>Title، H1 و Meta یکتا و هم‌راستا با نیت‌اند.</li><li>slug مطابق اکسل و canonical خودارجاع است.</li><li>صفحه index و داخل sitemap است.</li><li>تصویر، نام، alt و حجم صحیح‌اند.</li><li>یک لینک تجاری و دو لینک محتوایی مرتبط وجود دارد.</li><li>URL با ۲۰۰ باز می‌شود و redirect chain ندارد.</li></ol><h2>۷. منابع رسمی</h2><p><a href='https://developers.google.com/search/docs/crawling-indexing/url-structure'>Google Search Central — URL Structure</a><br><a href='https://developers.google.com/search/docs/appearance/google-images'>Google Search Central — Google Images SEO</a></p></div></body></html>''',encoding='utf-8')

if __name__=='__main__': workbook(); folders(); html(); print(BOOK); print(HTML); print(QUEUE)
